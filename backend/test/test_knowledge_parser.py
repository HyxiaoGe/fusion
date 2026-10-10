import io
import signal
import unittest
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from app.services.knowledge.parser import KnowledgeDocumentParser, KnowledgeParseError, parse_document_isolated


class KnowledgeDocumentParserTests(unittest.TestCase):
    def setUp(self):
        self.parser = KnowledgeDocumentParser()

    def test_text_parser_normalizes_deterministically_without_truncation(self):
        content = ("Ａ  B\r\n" + "长文本" * 3000).encode()

        first = self.parser.parse(content, mimetype="text/plain", filename="note.txt")
        second = self.parser.parse(content, mimetype="text/plain", filename="note.txt")

        self.assertEqual(first, second)
        self.assertTrue(first[0].text.startswith("A B\n"))
        self.assertGreater(len(first[0].text), 6000)

    def test_non_utf8_text_is_rejected(self):
        with self.assertRaisesRegex(KnowledgeParseError, "UTF-8") as raised:
            self.parser.parse(b"\xff\xfe", mimetype="text/plain", filename="note.txt")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_ENCODING_UNSUPPORTED")

    def test_nul_in_parsed_text_is_rejected_as_invalid_document(self):
        with self.assertRaisesRegex(KnowledgeParseError, "NUL") as raised:
            self.parser.parse(b"safe\x00unsafe", mimetype="text/plain", filename="note.txt")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_INVALID")

    def test_scanned_pdf_is_rejected_without_ocr(self):
        import PyPDF2

        stream = io.BytesIO()
        writer = PyPDF2.PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(stream)

        with self.assertRaises(KnowledgeParseError) as raised:
            self.parser.parse(stream.getvalue(), mimetype="application/pdf", filename="scan.pdf")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_SCANNED_DOCUMENT_UNSUPPORTED")

    def test_docx_zip_bomb_is_rejected_before_python_docx_parsing(self):
        archive = MagicMock()
        archive.__enter__.return_value.infolist.return_value = [
            SimpleNamespace(file_size=KnowledgeDocumentParser.MAX_DOCX_UNCOMPRESSED_BYTES + 1, flag_bits=0)
        ]
        with (
            patch("app.services.knowledge.parser.zipfile.ZipFile", return_value=archive),
            self.assertRaises(KnowledgeParseError) as raised,
        ):
            KnowledgeDocumentParser._validate_docx_archive(b"fake")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_TOO_LARGE")

    def test_legacy_doc_is_not_supported(self):
        with self.assertRaises(KnowledgeParseError) as raised:
            self.parser.parse(b"legacy", mimetype="application/msword", filename="legacy.doc")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_UNSUPPORTED")

    def test_mime_suffix_and_magic_mismatch_is_rejected(self):
        with self.assertRaises(KnowledgeParseError) as raised:
            self.parser.parse(b"%PDF-fake", mimetype="text/plain", filename="manual.pdf")

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_TYPE_MISMATCH")

    def test_isolated_parser_maps_child_crash_to_non_retryable_document_error(self):
        context = MagicMock()
        parent = MagicMock()
        child = MagicMock()
        process = MagicMock(exitcode=-signal.SIGKILL)
        parent.poll.return_value = True
        parent.recv.side_effect = EOFError
        process.is_alive.return_value = False
        context.Pipe.return_value = (parent, child)
        context.Process.return_value = process

        with (
            patch("app.services.knowledge.parser.multiprocessing.get_context", return_value=context),
            self.assertRaises(KnowledgeParseError) as raised,
        ):
            parse_document_isolated(
                b"%PDF-fake",
                mimetype="application/pdf",
                filename="manual.pdf",
                timeout_seconds=60,
            )

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_INVALID")
        process.join.assert_called()

    def test_isolated_parser_maps_cpu_limit_exit_to_parse_timeout(self):
        context = MagicMock()
        parent = MagicMock()
        child = MagicMock()
        process = MagicMock(exitcode=-signal.SIGXCPU)
        parent.poll.return_value = True
        parent.recv.side_effect = EOFError
        process.is_alive.return_value = False
        context.Pipe.return_value = (parent, child)
        context.Process.return_value = process

        with (
            patch("app.services.knowledge.parser.multiprocessing.get_context", return_value=context),
            self.assertRaises(KnowledgeParseError) as raised,
        ):
            parse_document_isolated(
                b"PK-fake",
                mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename="manual.docx",
                timeout_seconds=60,
            )

        self.assertEqual(raised.exception.code, "KNOWLEDGE_DOCUMENT_PARSE_TIMEOUT")


if __name__ == "__main__":
    unittest.main()


class StructuredKnowledgeParserTests(unittest.TestCase):
    def setUp(self):
        self.parser = KnowledgeDocumentParser()

    def test_markdown_sections_carry_heading_path_and_skip_fenced_headings(self):
        content = "\n".join(
            [
                "# 星槎 X3 手册",
                "",
                "## 售后",
                "### 延保",
                "两年延保 199 元。",
                "```bash",
                "# 不是标题",
                "```",
                "## 网络设置",
                "默认地址 192.168.77.1。",
            ]
        ).encode()

        sections = self.parser.parse(content, mimetype="text/markdown", filename="x3.md")

        self.assertEqual(
            [section.section for section in sections],
            ["星槎 X3 手册 > 售后 > 延保", "星槎 X3 手册 > 网络设置"],
        )
        self.assertTrue(sections[0].text.startswith("# 星槎 X3 手册"))
        self.assertIn("# 不是标题", sections[0].text)
        self.assertIn("192.168.77.1", sections[1].text)

    def test_markdown_without_headings_stays_single_unlabeled_section(self):
        sections = self.parser.parse("正文一\n\n正文二".encode(), mimetype="text/markdown", filename="note.md")

        self.assertEqual(len(sections), 1)
        self.assertIsNone(sections[0].section)

    def test_long_heading_path_keeps_most_specific_tail_within_label_limit(self):
        content = ("# " + "长" * 200 + "\n## 末级标题\n正文").encode()

        sections = self.parser.parse(content, mimetype="text/markdown", filename="note.md")

        self.assertEqual(len(sections[0].section), 120)
        self.assertTrue(sections[0].section.endswith(" > 末级标题"))

    def test_docx_groups_paragraphs_and_tables_under_headings_in_body_order(self):
        import docx

        document = docx.Document()
        document.add_heading("售后", level=1)
        document.add_paragraph("保修一年。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "套餐"
        table.cell(0, 1).text = "价格"
        table.cell(1, 0).text = "两年延保"
        table.cell(1, 1).text = "199 元"
        document.add_heading("延保", level=2)
        document.add_paragraph("线上购买。")
        document.add_heading("网络设置", level=1)
        document.add_paragraph("默认地址 192.168.77.1。")
        stream = io.BytesIO()
        document.save(stream)

        sections = self.parser.parse(
            stream.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            filename="x3.docx",
        )

        self.assertEqual([section.section for section in sections], ["售后", "售后 > 延保", "网络设置"])
        self.assertIn("保修一年。", sections[0].text)
        self.assertIn("两年延保 | 199 元", sections[0].text)
        self.assertIn("线上购买。", sections[1].text)
