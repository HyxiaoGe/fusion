"""知识库检索评测：在真实 Embedding 与 Milvus 上比较「解析版本 × 检索方式」的命中率。

语料与题目在 evals/knowledge_retrieval/。每个解析版本建一个临时 collection（schema v2，含 BM25），
同一份数据分别跑纯稠密检索与稠密 + BM25 混合检索；结束后删除临时 collection。

    python -m scripts.knowledge_retrieval_eval [--limit 24] [--show-misses]

需要与 Worker 相同的知识库 Embedding 与 Milvus 配置；不读写 PostgreSQL。
"""

from __future__ import annotations

import argparse
import asyncio
import io
import os
import unicodedata
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from app.ai.embeddings.base import EmbeddingProfile
from app.ai.embeddings.litellm_embedding import LiteLLMEmbeddingAdapter
from app.core.config import settings
from app.services.knowledge.chunker import DeterministicKnowledgeChunker, KnowledgeChunk, knowledge_index_text
from app.services.knowledge.milvus import KnowledgeVectorRecord, MilvusKnowledgeStore
from app.services.knowledge.parser import KnowledgeDocumentParser

EVAL_DIR = Path(__file__).resolve().parent.parent / "evals" / "knowledge_retrieval"
DOCX_MIMETYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
EVAL_USER_ID = "knowledge-eval"
EVAL_KNOWLEDGE_BASE_ID = "knowledge-eval-kb"
TOOL_TOP_K = 8


@dataclass
class EvalDocument:
    name: str
    filename: str
    mimetype: str
    content: bytes


@dataclass
class Variant:
    parser_version: str
    collection: str
    index_version: str
    chunks: dict[str, tuple[str, KnowledgeChunk]] = field(default_factory=dict)


def _build_docx(source: str) -> bytes:
    import docx

    document = docx.Document()
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, cells in enumerate(table_rows):
            for column_index, cell in enumerate(cells):
                table.cell(row_index, column_index).text = cell
        table_rows.clear()

    for line in source.splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            table_rows.append([cell.strip() for cell in stripped.strip("|").split("|")])
            continue
        flush_table()
        if not stripped:
            continue
        level = len(stripped) - len(stripped.lstrip("#"))
        if level:
            document.add_heading(stripped[level:].strip(), level=level)
        else:
            document.add_paragraph(stripped)
    flush_table()
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def load_documents() -> list[EvalDocument]:
    documents = []
    for path in sorted((EVAL_DIR / "documents").iterdir()):
        if path.name.startswith("."):
            continue
        if path.name.endswith(".docx.md"):
            documents.append(
                EvalDocument(
                    path.name,
                    path.name.removesuffix(".md"),
                    DOCX_MIMETYPE,
                    _build_docx(path.read_text(encoding="utf-8")),
                )
            )
        elif path.suffix == ".md":
            documents.append(EvalDocument(path.name, path.name, "text/markdown", path.read_bytes()))
    return documents


def load_questions() -> list[dict[str, Any]]:
    return yaml.safe_load((EVAL_DIR / "questions.yaml").read_text(encoding="utf-8"))["questions"]


def _comparable(text: str) -> str:
    # 表格单元格分隔符在 v1（制表符折叠为空格）与 v2（竖线）中不同，比较时一并忽略。
    return "".join(
        character for character in unicodedata.normalize("NFKC", text) if not character.isspace() and character != "|"
    )


def _profile() -> EmbeddingProfile:
    return EmbeddingProfile(
        settings.KNOWLEDGE_EMBEDDING_PROVIDER,
        settings.KNOWLEDGE_EMBEDDING_MODEL,
        settings.KNOWLEDGE_EMBEDDING_DIMENSION,
        settings.KNOWLEDGE_DISTANCE_METRIC,
        None,
        settings.KNOWLEDGE_EMBEDDING_REVISION,
    )


async def _embed(adapter: LiteLLMEmbeddingAdapter, profile: EmbeddingProfile, texts: list[str]) -> list[list[float]]:
    vectors: list[list[float]] = []
    batch_size = settings.KNOWLEDGE_EMBEDDING_BATCH_SIZE
    for offset in range(0, len(texts), batch_size):
        vectors.extend(await adapter.embed(texts[offset : offset + batch_size], profile))
    return vectors


async def populate_variant(
    client: Any,
    adapter: LiteLLMEmbeddingAdapter,
    profile: EmbeddingProfile,
    documents: list[EvalDocument],
    variant: Variant,
) -> None:
    parser_version = variant.parser_version
    parser = KnowledgeDocumentParser()
    chunker = DeterministicKnowledgeChunker(
        chunk_size=settings.KNOWLEDGE_CHUNK_SIZE,
        overlap=settings.KNOWLEDGE_CHUNK_OVERLAP,
    )
    structured = parser_version == KnowledgeDocumentParser.VERSION
    records: list[tuple[KnowledgeVectorRecord, str]] = []
    for document in documents:
        sections = parser.parse(
            document.content,
            mimetype=document.mimetype,
            filename=document.filename,
            version=parser_version,
        )
        for chunk in chunker.chunk(sections, document_id=document.name, index_version=variant.index_version):
            variant.chunks[chunk.chunk_id] = (document.name, chunk)
            # 与生产一致：v2 写入并嵌入「标题路径 + 正文」，v1 只用正文。
            index_text = knowledge_index_text(chunk) if structured else chunk.text
            record = KnowledgeVectorRecord(
                chunk=chunk,
                vector=[],
                user_id=EVAL_USER_ID,
                knowledge_base_id=EVAL_KNOWLEDGE_BASE_ID,
                document_id=document.name,
                index_version=variant.index_version,
                filename=document.filename,
            )
            records.append((record, index_text))
    vectors = await _embed(adapter, profile, [index_text for _record, index_text in records])
    MilvusKnowledgeStore._create_collection(
        client, variant.collection, profile.dimension, profile.distance_metric, hybrid=True
    )
    payload = []
    for (record, index_text), vector in zip(records, vectors, strict=True):
        row = MilvusKnowledgeStore._record_payload(record)
        row["vector"] = vector
        row["text"] = index_text
        payload.append(row)
    client.insert(collection_name=variant.collection, data=payload)
    client.flush(collection_name=variant.collection)
    client.load_collection(collection_name=variant.collection)


def _search(client: Any, variant: Variant, mode: str, query: str, vector: list[float], limit: int) -> list[str]:
    search_filter = MilvusKnowledgeStore.build_search_filter(
        EVAL_USER_ID, [EVAL_KNOWLEDGE_BASE_ID], [variant.index_version]
    )
    output_fields = ["document_id"]
    if mode == "dense":
        result = client.search(
            collection_name=variant.collection,
            data=[vector],
            anns_field="vector",
            filter=search_filter,
            limit=limit,
            output_fields=output_fields,
            search_params={"metric_type": settings.KNOWLEDGE_DISTANCE_METRIC},
            consistency_level="Strong",
        )
    else:
        result = MilvusKnowledgeStore._hybrid_search(
            client,
            collection=variant.collection,
            query_vector=vector,
            query_text=query,
            search_filter=search_filter,
            metric=settings.KNOWLEDGE_DISTANCE_METRIC,
            limit=limit,
            output_fields=output_fields,
        )
    rows = result[0] if result else []
    return [str(row.get("id") or row.get("chunk_id")) for row in rows]


def _rank(variant: Variant, chunk_ids: list[str], expected: dict[str, str]) -> int | None:
    answer = _comparable(expected["answer"])
    for rank, chunk_id in enumerate(chunk_ids, start=1):
        document_name, chunk = variant.chunks[chunk_id]
        if document_name == expected["document"] and answer in _comparable(chunk.text):
            return rank
    return None


def _summary(ranks: list[int | None]) -> str:
    total = len(ranks)

    def hit(k: int) -> float:
        return sum(1 for rank in ranks if rank is not None and rank <= k) / total

    mrr = sum(1 / rank for rank in ranks if rank is not None and rank <= TOOL_TOP_K) / total
    return f"hit@1 {hit(1):.2f}  hit@3 {hit(3):.2f}  hit@8 {hit(8):.2f}  MRR@8 {mrr:.3f}"


async def main() -> None:
    arguments = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    arguments.add_argument("--limit", type=int, default=TOOL_TOP_K * 3, help="每路召回条数，默认与工具 overfetch 一致")
    arguments.add_argument("--show-misses", action="store_true", help="列出 top 8 未命中的题目")
    arguments.add_argument("--show-ranks", action="store_true", help="最后输出每道题在各组合下的名次")
    options = arguments.parse_args()

    documents = load_documents()
    questions = load_questions()
    profile = _profile()
    adapter = LiteLLMEmbeddingAdapter()
    client = MilvusKnowledgeStore._build_client()
    run_id = os.environ.get("KNOWLEDGE_EVAL_RUN_ID") or uuid.uuid4().hex[:8]
    variants: list[Variant] = []
    try:
        for parser_version in (KnowledgeDocumentParser.LEGACY_VERSION, KnowledgeDocumentParser.VERSION):
            # 先登记再建库，建库中途失败也会在 finally 里清掉。
            variant = Variant(
                parser_version=parser_version,
                collection=f"zz_eval_{run_id}_{parser_version.replace('-', '_')}_v2_d{profile.dimension}",
                index_version=f"eval-{parser_version}",
            )
            variants.append(variant)
            await populate_variant(client, adapter, profile, documents, variant)
        query_vectors = await _embed(adapter, profile, [question["query"] for question in questions])

        print(f"语料 {len(documents)} 篇，题目 {len(questions)} 道，每路召回 {options.limit} 条")
        for variant in variants:
            print(f"{variant.parser_version}: {len(variant.chunks)} 个分块")
        kinds = sorted({question["kind"] for question in questions})
        rank_table: dict[str, list[int | None]] = {}
        for variant in variants:
            for mode in ("dense", "hybrid"):
                ranks = [
                    _rank(
                        variant,
                        _search(client, variant, mode, question["query"], vector, options.limit),
                        question["expected"],
                    )
                    for question, vector in zip(questions, query_vectors, strict=True)
                ]
                rank_table[f"{variant.parser_version}+{mode}"] = ranks
                print(f"\n[{variant.parser_version} + {mode}] 全部 {_summary(ranks)}")
                for kind in kinds:
                    kind_ranks = [rank for rank, question in zip(ranks, questions) if question["kind"] == kind]
                    print(f"  {kind:<10} ({len(kind_ranks):>2}) {_summary(kind_ranks)}")
                if options.show_misses:
                    for rank, question in zip(ranks, questions, strict=True):
                        if rank is None or rank > TOOL_TOP_K:
                            print(f"    未进前 {TOOL_TOP_K}: {question['id']}（名次 {rank or '-'}）")
        if options.show_ranks:
            print("\n" + "\t".join(["题目", *rank_table]))
            for index, question in enumerate(questions):
                print("\t".join([question["id"], *(str(ranks[index] or "-") for ranks in rank_table.values())]))
    finally:
        for variant in variants:
            if client.has_collection(collection_name=variant.collection):
                client.drop_collection(collection_name=variant.collection)
        client.close()


if __name__ == "__main__":
    asyncio.run(main())
